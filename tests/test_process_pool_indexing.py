from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from local_full_text_search.config.defaults import AppSettings
from local_full_text_search.core.database import DatabaseManager
from local_full_text_search.core.index_manager import IndexManager
from local_full_text_search.core.search_engine import SearchEngine
from local_full_text_search.models.search_query import SearchQuery


class ProcessPoolIndexingTests(unittest.TestCase):
    def test_docx_and_xlsx_use_process_lane_and_clean_spool_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            root = base / "files"
            root.mkdir()

            document = Document()
            document.add_paragraph("PROCESS_DOCX_TEST_HIT")
            document.save(root / "sample.docx")

            workbook = Workbook()
            workbook.active["A1"] = "PROCESS_XLSX_TEST_HIT"
            workbook.save(root / "sample.xlsx")

            db = DatabaseManager(base / "index.db")
            db.initialize()
            root_id = db.add_root(root)
            queues: list[str] = []
            worker_pids: list[int] = []
            descriptor_bytes: list[int] = []

            def capture_progress(payload: dict[str, object]) -> None:
                if payload.get("stage") == "indexing" and payload.get("completed_queue"):
                    queues.append(str(payload["completed_queue"]))
                    if isinstance(payload.get("worker_pid"), int):
                        worker_pids.append(int(payload["worker_pid"]))
                    if isinstance(payload.get("process_descriptor_bytes"), int):
                        descriptor_bytes.append(int(payload["process_descriptor_bytes"]))

            settings = AppSettings(
                enable_ocr=False,
                process_parser_workers=1,
                process_pending_tasks=1,
                process_max_tasks_per_child=1,
                large_office_process_min_bytes=0,
            )
            temp_root = base / "runtime_temp"
            with patch("local_full_text_search.core.index_manager.TEMP_DIR", temp_root):
                summary = IndexManager(db, settings).index_root(root_id, progress_callback=capture_progress)

            engine = SearchEngine(db)
            self.assertEqual(summary.indexed, 2)
            self.assertEqual(summary.failed, 0)
            self.assertEqual(queues.count("office_process"), 2)
            self.assertEqual(len(set(worker_pids)), 2)
            self.assertTrue(descriptor_bytes)
            self.assertLess(max(descriptor_bytes), 4096)
            self.assertEqual(engine.search(SearchQuery(text="PROCESS_DOCX_TEST_HIT")).total_confirmed, 1)
            self.assertEqual(engine.search(SearchQuery(text="PROCESS_XLSX_TEST_HIT")).total_confirmed, 1)
            spool_parent = temp_root / "process_results"
            self.assertFalse(spool_parent.exists() and any(spool_parent.iterdir()))


if __name__ == "__main__":
    unittest.main()
