from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest

from local_full_text_search.core.compact_snapshot_builder import _code
from local_full_text_search.core.block_coalescer import BlockCoalescer
from local_full_text_search.core.database import DatabaseManager
from local_full_text_search.core.index_generation_store import IndexGenerationStore
from local_full_text_search.core.task_manager import CancelToken
from local_full_text_search.models.content_block import ContentBlock
from local_full_text_search.parsers.docx_parser import DocxParser
from local_full_text_search.parsers.pdf_parser import PdfParser
from local_full_text_search.parsers.text_parser import TextParser
from local_full_text_search.rag.block_reader import BlockReader, DocumentBlocks
from local_full_text_search.rag.chunk_builder import ChunkBuilder, ChunkPolicy
from local_full_text_search.rag.models import (
    FileSource,
    NoPublishedGenerationError,
    RagChunk,
    SourceBlock,
    SourceRef,
    SourceSpan,
)


def _publish_fixture(
    tmp_path: Path,
    sources: list[tuple[str, str, list[ContentBlock]]],
) -> tuple[Path, int]:
    base = tmp_path / "search_index.db"
    store = IndexGenerationStore(base)
    candidate = store.begin_candidate()
    db = DatabaseManager(candidate.path)
    db.initialize()
    root = tmp_path / "indexed"
    root.mkdir()
    root_id = db.add_root(root)
    block_count = 0
    for filename, parser_name, blocks in sources:
        file_path = root / filename
        file_path.write_text("fixture", encoding="utf-8")
        file_id, _changed = db.upsert_file_metadata(root_id, file_path)
        db.replace_file_blocks(
            file_id,
            file_path.name,
            str(file_path),
            blocks,
            parser_name=parser_name,
        )
        block_count += len(blocks)
    published = store.publish(
        candidate,
        file_count=len(sources),
        document_count=len(sources),
        block_count=block_count,
        published_at="2026-10-15T00:00:00+00:00",
    )
    return base, published.generation


def _block(
    path: Path,
    index: int,
    kind: str,
    text: str,
    *,
    location: str = "",
    page: int | None = None,
    line_start: int | None = None,
    line_end: int | None = None,
    source_type: str = "native_text",
    confidence: float | None = None,
) -> ContentBlock:
    return ContentBlock(
        file_path=str(path),
        block_index=index,
        block_type=kind,
        location_text=location,
        raw_text=text,
        normalized_text=text.lower(),
        page_number=page,
        line_start=line_start,
        line_end=line_end,
        source_type=source_type,
        ocr_confidence=confidence,
    )


def test_models_reject_invalid_source_ranges() -> None:
    with pytest.raises(ValueError, match="half-open"):
        SourceSpan(block_id=1, block_index=0, raw_start=3, raw_end=3)

    source = SourceRef(
        generation=1,
        stable_chunk_id="chunk-a",
        file_id=1,
        root_id=1,
        document_id=1,
        file_path="C:/docs/a.txt",
        filename="a.txt",
        extension=".txt",
        spans=(SourceSpan(1, 0, 0, 1),),
        block_index_start=0,
        block_index_end=0,
        location_text="line 1",
    )
    with pytest.raises(ValueError, match="refer to this stable_chunk_id"):
        RagChunk("chunk-b", "a", "a", "hash", "evidence-v1", (source,))


def test_reader_uses_only_active_published_compact_generation(tmp_path: Path) -> None:
    path = tmp_path / "source.pdf"
    base, generation = _publish_fixture(
        tmp_path,
        [
            (
                path.name,
                "pdf",
                [
                    _block(path, 0, "pdf_page", "native page text", location="第 1 页", page=1),
                    _block(
                        path,
                        1,
                        "pdf_page_ocr",
                        "scanned page text",
                        location="第 2 页 OCR",
                        page=2,
                        source_type="ocr",
                        confidence=0.73,
                    ),
                ],
            )
        ],
    )
    store = IndexGenerationStore(base)
    building = store.begin_candidate()
    try:
        with BlockReader(base).open_active() as reader:
            docs = list(reader.iter_documents())
            assert reader.generation == generation
            assert docs[0].generation == generation
            assert [block.block_kind for block in docs[0].blocks] == [
                "pdf_page",
                "pdf_page_ocr",
            ]
            assert docs[0].blocks[1].ocr_confidence == pytest.approx(0.73)
            with pytest.raises(sqlite3.OperationalError):
                reader.connection.execute("UPDATE files SET filename = 'changed'")
    finally:
        store.abort_candidate(building)


def test_reader_errors_when_only_building_database_exists(tmp_path: Path) -> None:
    store = IndexGenerationStore(tmp_path / "search_index.db")
    candidate = store.begin_candidate()
    try:
        with pytest.raises(NoPublishedGenerationError):
            with BlockReader(store.database_path).open_active():
                pass
    finally:
        store.abort_candidate(candidate)


def test_reader_filters_formats_and_returns_block_locations(tmp_path: Path) -> None:
    txt = tmp_path / "notes.txt"
    pdf = tmp_path / "paper.pdf"
    base, _generation = _publish_fixture(
        tmp_path,
        [
            (txt.name, "text", [_block(txt, 0, "text", "line one\nline two", location="第 1-2 行", line_start=1, line_end=2)]),
            (pdf.name, "pdf", [_block(pdf, 0, "pdf_page", "page", location="第 1 页", page=1)]),
        ],
    )
    with BlockReader(base).open_active() as reader:
        documents = list(reader.iter_documents(file_ids={1}, extensions=frozenset({".txt"})))
    assert len(documents) == 1
    assert documents[0].files[0].extension == ".txt"
    assert documents[0].blocks[0].line_start == 1
    assert documents[0].blocks[0].line_end == 2


def _document(
    *,
    generation: int,
    blocks: tuple[SourceBlock, ...],
    extension: str,
    document_id: int = 19,
    file_id: int = 5,
    path: str = "C:/docs/sample.txt",
) -> DocumentBlocks:
    return DocumentBlocks(
        generation=generation,
        document_id=document_id,
        files=(FileSource(file_id, 1, path, Path(path).name, extension, "text", "1"),),
        blocks=blocks,
    )


def test_pdf_chunks_never_cross_pages_and_keep_ocr_confidence() -> None:
    blocks = (
        SourceBlock(1, 19, 11, 0, "pdf_page", "甲" * 30, "第 1 页", page_number=1),
        SourceBlock(1, 19, 12, 1, "pdf_page_ocr", "乙" * 30, "第 2 页 OCR", page_number=2, source_type="ocr", ocr_confidence=0.6),
    )
    chunks = ChunkBuilder(ChunkPolicy(target_chars=15, max_chars=20, overlap_chars=3)).build(
        _document(generation=1, blocks=blocks, extension=".pdf")
    )
    assert chunks
    assert all(len(chunk.text) <= 20 for chunk in chunks)
    assert {chunk.sources[0].page_number for chunk in chunks} == {1, 2}
    assert {chunk.sources[0].source_type for chunk in chunks} == {"native_text", "ocr"}
    ocr_chunks = [chunk for chunk in chunks if chunk.sources[0].source_type == "ocr"]
    assert all(chunk.sources[0].ocr_confidence == pytest.approx(0.6) for chunk in ocr_chunks)


def test_txt_chunks_preserve_line_ranges_and_reconstruct_from_spans() -> None:
    raw = "first line\nsecond line\nthird line\nfourth line"
    block = SourceBlock(
        1,
        19,
        21,
        0,
        "text",
        raw,
        "第 1-4 行",
        line_start=1,
        line_end=4,
    )
    chunks = ChunkBuilder(ChunkPolicy(target_chars=18, max_chars=24, overlap_chars=4)).build(
        _document(generation=1, blocks=(block,), extension=".txt")
    )
    assert len(chunks) > 1
    for chunk in chunks:
        source = chunk.sources[0]
        reconstructed = "".join(
            span.separator_before + raw[span.raw_start : span.raw_end]
            for span in source.spans
        )
        assert reconstructed == chunk.text
        assert 1 <= source.line_start <= source.line_end <= 4


def test_docx_does_not_mix_paragraphs_and_table_rows_or_tables() -> None:
    blocks = (
        SourceBlock(1, 19, 31, 0, "docx_paragraph", "paragraph", "正文第 1 段 - 正文第 1 段"),
        SourceBlock(1, 19, 32, 1, "docx_table_row", "table one", "表格 1，第 1 行"),
        SourceBlock(1, 19, 33, 2, "docx_table_row", "table two", "表格 2，第 1 行"),
    )
    chunks = ChunkBuilder().build(_document(generation=1, blocks=blocks, extension=".docx"))
    assert [chunk.text for chunk in chunks] == ["paragraph", "table one", "table two"]


def test_stable_chunk_id_ignores_generation_and_database_row_ids() -> None:
    first_block = SourceBlock(1, 19, 41, 0, "text", "same text", "第 1 行", line_start=1, line_end=1)
    second_block = SourceBlock(8, 28, 991, 0, "text", "same text", "第 1 行", line_start=1, line_end=1)
    builder = ChunkBuilder()
    first = builder.build(_document(generation=1, blocks=(first_block,), extension=".txt"))[0]
    second = builder.build(
        _document(
            generation=8,
            blocks=(second_block,),
            extension=".txt",
            document_id=28,
            file_id=88,
            path="D:/moved/sample.txt",
        )
    )[0]
    assert first.stable_chunk_id == second.stable_chunk_id
    assert first.sources[0].generation == 1
    assert second.sources[0].generation == 8


def test_duplicate_document_keeps_multiple_file_sources() -> None:
    block = SourceBlock(1, 19, 61, 0, "text", "shared document", "第 1 行", line_start=1, line_end=1)
    document = DocumentBlocks(
        generation=1,
        document_id=19,
        files=(
            FileSource(5, 1, "C:/docs/a.txt", "a.txt", ".txt", "text", "1"),
            FileSource(6, 2, "D:/archive/a.txt", "a.txt", ".txt", "text", "1"),
        ),
        blocks=(block,),
    )
    chunks = ChunkBuilder().build(document)
    assert len(chunks) == 1
    assert [source.file_id for source in chunks[0].sources] == [5, 6]
    assert [source.file_path for source in chunks[0].sources] == [
        "C:/docs/a.txt",
        "D:/archive/a.txt",
    ]


def test_txt_chunk_spans_reconstruct_across_adjacent_database_blocks() -> None:
    first = SourceBlock(1, 19, 71, 0, "text", "one\ntwo", "第 1-2 行", line_start=1, line_end=2)
    second = SourceBlock(1, 19, 72, 1, "text", "three\nfour", "第 3-4 行", line_start=3, line_end=4)
    document = _document(generation=1, blocks=(first, second), extension=".txt")
    chunks = ChunkBuilder(ChunkPolicy(target_chars=100, max_chars=120, overlap_chars=0)).build(document)
    assert len(chunks) == 1
    chunk = chunks[0]
    source = chunk.sources[0]
    by_id = {first.block_id: first, second.block_id: second}
    reconstructed = "".join(
        span.separator_before + by_id[span.block_id].raw_text[span.raw_start : span.raw_end]
        for span in source.spans
    )
    assert reconstructed == chunk.text == "one\ntwo\nthree\nfour"
    assert (source.line_start, source.line_end) == (1, 4)


def test_text_or_policy_change_changes_stable_id() -> None:
    original = SourceBlock(1, 19, 51, 0, "text", "old text", "第 1 行", line_start=1, line_end=1)
    changed = SourceBlock(1, 19, 52, 0, "text", "new text", "第 1 行", line_start=1, line_end=1)
    a = ChunkBuilder().build(_document(generation=1, blocks=(original,), extension=".txt"))[0]
    b = ChunkBuilder().build(_document(generation=1, blocks=(changed,), extension=".txt"))[0]
    c = ChunkBuilder(ChunkPolicy(version="evidence-v2")).build(
        _document(generation=1, blocks=(original,), extension=".txt")
    )[0]
    assert len({a.stable_chunk_id, b.stable_chunk_id, c.stable_chunk_id}) == 3


def test_chunk_builder_enforces_embedding_token_limit_without_losing_source_offsets() -> None:
    raw = "第一段内容。第二段内容。第三段内容。第四段内容。"
    block = SourceBlock(1, 19, 81, 0, "text", raw, "正文", line_start=1, line_end=1)
    builder = ChunkBuilder(
        ChunkPolicy(
            version="evidence-v2",
            target_chars=16,
            max_chars=24,
            overlap_chars=2,
            embedding_model_revision="test-revision",
            max_embedding_tokens=8,
        ),
        token_counter=len,
    )
    chunks = builder.build(_document(generation=1, blocks=(block,), extension=".txt"))
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 8 for chunk in chunks)
    assert all(chunk.chunk_policy_version == "evidence-v2" for chunk in chunks)
    for chunk in chunks:
        source = chunk.sources[0]
        rebuilt = "".join(
            span.separator_before + raw[span.raw_start : span.raw_end]
            for span in source.spans
        )
        assert rebuilt == chunk.text


def test_compact_type_codes_match_snapshot_builder() -> None:
    from local_full_text_search.rag.block_reader import _compact_kind_code

    for kind in ("pdf_page", "pdf_page_ocr", "docx_paragraph", "docx_table_row", "text"):
        assert _compact_kind_code(kind) == _code(kind)


def test_real_pdf_docx_and_txt_parser_blocks_feed_chunk_builder(tmp_path: Path) -> None:
    import fitz
    from docx import Document as WordDocument

    token = CancelToken()

    pdf_path = tmp_path / "parser.pdf"
    pdf = fitz.open()
    for page_number in range(2):
        page = pdf.new_page()
        page.insert_text((72, 72), f"PDF page {page_number + 1} evidence")
    pdf.save(pdf_path)
    pdf.close()
    pdf_blocks = BlockCoalescer().coalesce(list(PdfParser().parse(pdf_path, token)))
    assert [block.block_type for block in pdf_blocks] == ["pdf_page", "pdf_page"]
    pdf_document = DocumentBlocks(
        generation=1,
        document_id=101,
        files=(FileSource(101, 1, str(pdf_path), pdf_path.name, ".pdf", "pdf", "1"),),
        blocks=tuple(
            SourceBlock(
                1, 101, index + 1, block.block_index, block.block_type,
                block.raw_text, block.location_text, page_number=block.page_number,
                source_type=block.source_type,
            )
            for index, block in enumerate(pdf_blocks)
        ),
    )
    pdf_chunks = ChunkBuilder().build(pdf_document)
    assert [chunk.sources[0].page_number for chunk in pdf_chunks] == [1, 2]

    docx_path = tmp_path / "parser.docx"
    word = WordDocument()
    word.add_paragraph("First paragraph has evidence.")
    word.add_paragraph("Second paragraph has more evidence.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Header"
    table.cell(0, 1).text = "Value"
    word.save(docx_path)
    docx_blocks = BlockCoalescer().coalesce(list(DocxParser().parse(docx_path, token)))
    assert {block.block_type for block in docx_blocks} == {"docx_paragraph", "docx_table_row"}
    docx_document = DocumentBlocks(
        generation=1,
        document_id=102,
        files=(FileSource(102, 1, str(docx_path), docx_path.name, ".docx", "docx", "1"),),
        blocks=tuple(
            SourceBlock(
                1, 102, index + 1, block.block_index, block.block_type,
                block.raw_text, block.location_text,
            )
            for index, block in enumerate(docx_blocks)
        ),
    )
    docx_chunks = ChunkBuilder().build(docx_document)
    assert len(docx_chunks) == 2
    assert all("表格" not in chunk.sources[0].location_text for chunk in docx_chunks[:1])
    assert "表格" in docx_chunks[-1].sources[0].location_text

    txt_path = tmp_path / "parser.txt"
    txt_path.write_text("line one\nline two\nline three\n", encoding="utf-8")
    txt_blocks = BlockCoalescer().coalesce(list(TextParser(lines_per_block=1).parse(txt_path, token)))
    assert len(txt_blocks) == 1
    assert (txt_blocks[0].line_start, txt_blocks[0].line_end) == (1, 3)
    txt_document = DocumentBlocks(
        generation=1,
        document_id=103,
        files=(FileSource(103, 1, str(txt_path), txt_path.name, ".txt", "text", "1"),),
        blocks=(
            SourceBlock(
                1, 103, 103, txt_blocks[0].block_index, txt_blocks[0].block_type,
                txt_blocks[0].raw_text, txt_blocks[0].location_text,
                line_start=txt_blocks[0].line_start, line_end=txt_blocks[0].line_end,
            ),
        ),
    )
    assert ChunkBuilder().build(txt_document)[0].sources[0].line_start == 1
